"""Build the IEEE two-column manuscript from MBCA_paper.tex.

    python paper/build_pdf.py            # PDF (LaTeX) and DOCX
    python paper/build_pdf.py --pdf      # PDF only
    python paper/build_pdf.py --docx     # DOCX only

PDF: the real IEEEtran class, compiled with Tectonic (`tectonic -X compile`, BibTeX with IEEEtran.bst). Tectonic is
found on PATH, via $TECTONIC, or in %LOCALAPPDATA%/Programs/tectonic. The same .tex compiles unchanged on Overleaf.

DOCX: pandoc reads a flattened copy of the .tex -- every \\input resolved, the `latex-only` preamble block dropped,
`table*` turned into `table` (pandoc's LaTeX reader does not parse `table*`) and IEEEkeywords into an "Index Terms"
paragraph -- against an IEEE-styled reference document. A python-docx pass then applies the IEEE look: Letter page,
two 3.487 in columns, Times New Roman, I/A section numbers, "Fig. 1." and "TABLE I" captions, cross-references in
the same form, booktabs table rules, and full-width sections for wide figures and tables.

Requires pypandoc-binary and python-docx (and Tectonic for the PDF).
"""
import os
import re
import shutil
import subprocess
import sys
from copy import deepcopy
from pathlib import Path

PAPER = Path(__file__).resolve().parent
TEX_SRC = PAPER / "MBCA_paper.tex"
BIB = PAPER / "references.bib"
CSL = PAPER / "assets" / "ieee.csl"
REF_DOCX = PAPER / "assets" / "ieee_ref.docx"
OUT_DOCX = PAPER / "MBCA_paper.docx"
OUT_PDF = PAPER / "MBCA_paper.pdf"
FLAT = PAPER / ".pandoc_flat.tex"

# IEEEtran journal geometry (US Letter): 43 pc text width, two 21 pc columns, 1 pc gutter.
TEXT_W, COL_W, TOP, BOTTOM = 7.14, 3.487, 0.75, 1.0
SIDE = (8.5 - TEXT_W) / 2
UNNUMBERED = {"References", "Data and Code Availability"}
CHAR_W, CELL_PAD = 0.056, 0.17  # inches: mean 8 pt Times character; Word's default left+right cell margins
FONT = "Times New Roman"


# =========================================================================== PDF
def tectonic() -> str:
    for cand in (os.environ.get("TECTONIC"), shutil.which("tectonic"),
                 Path(os.environ.get("LOCALAPPDATA", "")) / "Programs" / "tectonic" / "tectonic.exe"):
        if cand and Path(cand).exists():
            return str(cand)
    sys.exit("Tectonic not found: install it (https://tectonic-typesetting.github.io) or set $TECTONIC")


def build_pdf() -> None:
    r = subprocess.run([tectonic(), "-X", "compile", TEX_SRC.name], cwd=PAPER, capture_output=True, text=True)
    log = r.stdout + r.stderr
    problems = [ln for ln in log.splitlines()
                if re.search(r"undefined|multiply defined|Overfull \\hbox \((?:[2-9]|\d\d)", ln, re.I)]
    if r.returncode:
        sys.exit(log[-3000:])
    print("\n".join(problems) or "  LaTeX: no undefined references, no overfull boxes > 2pt")
    print(f"done: {OUT_PDF}")


# =========================================================================== DOCX: source preparation
def flatten() -> list[bool]:
    """Write the pandoc copy of the manuscript; return, in order, whether each table is full width (table*)."""
    def resolve(text: str) -> str:
        return re.sub(r"\\input\{([^}]+)\}", lambda m: resolve((PAPER / m.group(1)).read_text(encoding="utf-8")), text)

    src = resolve(TEX_SRC.read_text(encoding="utf-8"))
    src = re.sub(r"% --- latex-only begin.*?% --- latex-only end", "", src, flags=re.S)
    wide = [m.group(1) == "*" for m in re.finditer(r"\\begin\{table(\*?)\}", src)]
    src = src.replace("\\begin{table*}", "\\begin{table}").replace("\\end{table*}", "\\end{table}")
    src = re.sub(r"\\begin\{IEEEkeywords\}\s*(.*?)\s*\\end\{IEEEkeywords\}", r"Index Terms---\1", src, flags=re.S)
    src = src.replace("\\appendices", "")
    FLAT.write_text(src, encoding="utf-8")
    return wide


def restyle_reference_doc() -> None:
    """IEEE styles on the handful of styles pandoc emits (idempotent; the reference doc is a committed asset)."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document(str(REF_DOCX))

    def style(name, size, bold=False, italic=False, caps=False, align=None, before=0, after=0, indent=None,
              font=FONT, keep=False):
        try:
            st = doc.styles[name]
        except KeyError:
            return
        f = st.font
        f.name, f.size, f.bold, f.italic, f.small_caps = font, size and Pt(size), bold, italic, caps
        f.color.rgb = RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr()
        fonts = rpr.find(qn("w:rFonts"))
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            rpr.append(fonts)
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            fonts.set(qn(a), font)
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):  # theme fonts override the name
            fonts.attrib.pop(qn(a), None)
        if st.type != 1:  # character style: font only
            return
        pf = st.paragraph_format
        pf.space_before, pf.space_after, pf.line_spacing = Pt(before), Pt(after), 1.0
        if align is not None:
            pf.alignment = align
        pf.first_line_indent = indent if indent is not None else Inches(0)  # never inherit the body indent
        pf.keep_with_next = keep

    body_indent = Inches(0.125)
    style("Normal", 10, align=AL.JUSTIFY)
    style("Body Text", 10, align=AL.JUSTIFY, indent=body_indent)
    style("First Paragraph", 10, align=AL.JUSTIFY, indent=body_indent)
    style("Compact", 8, align=AL.LEFT)
    style("Title", 24, align=AL.CENTER, after=10)
    style("Author", 11, align=AL.CENTER, after=12)
    style("Abstract Title", 9, bold=True)
    style("Abstract", 9, bold=True, align=AL.JUSTIFY, after=4)
    style("Heading 1", 10, caps=True, align=AL.CENTER, before=10, after=4, keep=True)
    style("Heading 2", 10, italic=True, align=AL.LEFT, before=6, after=3, keep=True)
    style("Heading 3", 10, italic=True, align=AL.LEFT, before=4, after=2, keep=True)
    style("Table Caption", 8, caps=True, align=AL.CENTER, before=6, after=4, keep=True)
    style("Image Caption", 8, align=AL.JUSTIFY, before=2, after=8)
    style("Captioned Figure", 8, align=AL.CENTER, before=4, keep=True)
    style("Figure", 8, align=AL.CENTER, keep=True)
    style("Bibliography", 8, align=AL.LEFT, after=2)
    style("Footnote Text", 8, align=AL.JUSTIFY)
    style("Hyperlink", None)  # inherit the size of the surrounding text (captions, tables)
    style("Verbatim Char", 9, font="Courier New")
    doc.styles["Hyperlink"].font.underline = False
    doc.save(str(REF_DOCX))


def convert() -> None:
    import pypandoc

    pypandoc.convert_file(
        str(FLAT), to="docx", format="latex", outputfile=str(OUT_DOCX), cworkdir=str(PAPER),
        extra_args=[f"--reference-doc={REF_DOCX}", "--citeproc", f"--bibliography={BIB}", f"--csl={CSL}",
                    f"--resource-path={PAPER}", "--mathml", "-M", "reference-section-title=References"])


# =========================================================================== DOCX: IEEE post-pass
def roman(n: int) -> str:
    out = ""
    for v, s in ((1000, "M"), (900, "CM"), (500, "D"), (400, "CD"), (100, "C"), (90, "XC"), (50, "L"), (40, "XL"),
                 (10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")):
        while n >= v:
            out, n = out + s, n - v
    return out


def _w(tag):
    from docx.oxml.ns import qn
    return qn(f"w:{tag}")


def _set(el, **attrs):
    for k, v in attrs.items():
        el.set(_w(k), v)
    return el


def _style(el):
    ppr = el.find(_w("pPr"))
    ps = ppr.find(_w("pStyle")) if ppr is not None else None
    return ps.get(_w("val")) if ps is not None else None


def _text(el):
    return "".join(t.text or "" for t in el.iter(_w("t")))


def _prefix(p, text, italic=None, bold=None, br=False):
    """Insert a run carrying `text` at the start of paragraph element `p` (after its pPr)."""
    from docx.oxml import OxmlElement
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    for flag, tag in ((bold, "b"), (italic, "i")):
        if flag:
            rpr.append(OxmlElement(f"w:{tag}"))
    r.append(rpr)
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    if br:
        r.append(OxmlElement("w:br"))
    ppr = p.find(_w("pPr"))
    (ppr.addnext if ppr is not None else lambda x: p.insert(0, x))(r)


def number_and_crossref(doc) -> dict:
    """IEEE numbering of headings (I, A), figures (1) and tables (I), and every internal reference to them."""
    body = doc.element.body
    labels, pending = {}, []
    h1 = h2 = fig = tab = app = 0
    for el in list(body):
        if el.tag == _w("bookmarkStart"):
            name = el.get(_w("name"))
            if name and not name.startswith("ref"):
                pending.append(name)
            continue
        st = _style(el) if el.tag == _w("p") else None
        if st == "Heading1":
            if _text(el).strip() in UNNUMBERED:
                pending = []
                continue
            if any(n.startswith("app:") for n in pending):
                app += 1
                num = chr(64 + app)
                _prefix(el, f"Appendix {num}: ")
            else:
                h1 += 1
                num = roman(h1)
                _prefix(el, f"{num}. ")
            h2 = 0
            labels.update({n: num for n in pending})
        elif st == "Heading2":
            h2 += 1
            _prefix(el, f"{chr(64 + h2)}. ")
            labels.update({n: f"{roman(h1)}-{chr(64 + h2)}" for n in pending})
        elif st == "TableCaption":
            tab += 1
            _prefix(el, f"TABLE {roman(tab)}", br=True)
            labels.update({n: roman(tab) for n in pending})
        elif st == "CaptionedFigure":
            fig += 1
            labels.update({n: str(fig) for n in pending})
            cap = el.getnext()
            if cap is not None and _style(cap) == "ImageCaption":
                _prefix(cap, f"Fig. {fig}. ")
        elif el.tag not in (_w("p"), _w("tbl")):
            continue
        pending = []
    missing = set()
    for link in body.iter(_w("hyperlink")):
        anchor = link.get(_w("anchor"))
        if anchor is None or anchor.startswith("ref-") or anchor == "refs":
            continue
        if anchor not in labels:
            missing.add(anchor)
            continue
        ts = list(link.iter(_w("t")))
        if ts:
            ts[0].text = labels[anchor]
            for t in ts[1:]:
                t.text = ""
    if missing:
        print(f"  WARNING: references without a numbered target: {sorted(missing)}")
    print(f"  numbered {h1} sections, {app} appendices, {fig} figures, {tab} tables")
    return labels


def front_matter(doc) -> None:
    """IEEE abstract and index terms: bold 9 pt run-in paragraphs; drop pandoc's separate 'Abstract' title."""
    from docx.shared import Pt
    body = doc.element.body
    for el in list(body):
        st = _style(el) if el.tag == _w("p") else None
        if st == "AbstractTitle":
            body.remove(el)
        elif st == "Abstract":
            _prefix(el, "Abstract\u2014", italic=True, bold=True)
            break
    for p in doc.paragraphs:
        if p.text.startswith("Index Terms"):
            rest = "Index Terms\u2014"  # strip it run by run (pandoc may split the text across runs)
            for r in p.runs:
                k = len(os.path.commonprefix([r.text, rest]))
                r.text, rest = r.text[k:], rest[k:]
                if not rest or k == 0:
                    break
            _prefix(p._p, "Index Terms\u2014", italic=True, bold=True)
            for r in p.runs:
                r.font.bold, r.font.size = True, Pt(9)
            p.paragraph_format.first_line_indent = 0
            break


def booktabs(doc) -> None:
    """Top/bottom rules and a rule under the header row; no vertical lines; centered; header repeats."""
    from docx.oxml import OxmlElement
    for tbl in doc.element.body.iter(_w("tbl")):
        tblpr = tbl.find(_w("tblPr"))
        for tag in ("tblBorders", "jc"):
            old = tblpr.find(_w(tag))
            if old is not None:
                tblpr.remove(old)
        jc = OxmlElement("w:jc")
        jc.set(_w("val"), "center")
        borders = OxmlElement("w:tblBorders")
        for side, sz in (("top", "8"), ("left", None), ("bottom", "8"), ("right", None), ("insideH", None),
                         ("insideV", None)):
            b = _set(OxmlElement(f"w:{side}"), val="single" if sz else "nil")
            if sz:
                _set(b, sz=sz, space="0", color="000000")
            borders.append(b)
        tblpr.append(jc)
        tblpr.append(borders)
        rows = tbl.findall(_w("tr"))
        for tr in rows:
            trpr = tr.find(_w("trPr"))
            if trpr is None:
                trpr = OxmlElement("w:trPr")
                tr.insert(0, trpr)
            trpr.append(OxmlElement("w:cantSplit"))
            if trpr.find(_w("tblHeader")) is None:
                continue
            for tc in tr.findall(_w("tc")):
                tcpr = tc.find(_w("tcPr"))
                tcb = OxmlElement("w:tcBorders")
                tcb.append(_set(OxmlElement("w:bottom"), val="single", sz="4", space="0", color="000000"))
                tcpr.append(tcb)


TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW",
               "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]
TRPR_ORDER = ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter", "cantSplit", "trHeight", "tblHeader",
              "tblCellSpacing", "jc", "hidden"]


def _schema_order(parent, order):
    """Re-sort children into the OOXML schema sequence (Word rejects some out-of-order property elements)."""
    rank = {_w(t): i for i, t in enumerate(order)}
    kids = sorted(parent, key=lambda el: rank.get(el.tag, len(order)))
    for el in kids:
        parent.append(el)


def size_tables(doc, wide_tables: list[bool]) -> None:
    """Content-proportional column widths (pandoc gives plain l/r columns equal widths, which wraps numbers), sized to
    the column or the full text width, and each table kept on one page."""
    from docx.oxml import OxmlElement
    for tbl, wide in zip(doc.element.body.iter(_w("tbl")), wide_tables):
        rows = [[_text(tc) for tc in tr.findall(_w("tc"))] for tr in tbl.findall(_w("tr"))]
        ncol = max(len(r) for r in rows)
        # characters each column needs: its longest body cell, or the longest header *word* (headers may wrap)
        need = [max([len(r[j]) for r in rows[1:] if j < len(r)] + [len(w) for w in rows[0][j].split()] + [1])
                for j in range(ncol)]
        text = [n * CHAR_W for n in need]
        avail = (TEXT_W if wide else COL_W) - ncol * CELL_PAD
        short = [t if n <= 14 else 0.0 for t, n in zip(text, need)]
        if sum(text) <= avail:  # everything fits on one line: share the slack in proportion
            text = [t * avail / sum(text) for t in text]
        elif sum(short) < avail:  # short columns keep one line; long text columns wrap in the rest
            rest = [t for t, s in zip(text, short) if not s]
            text = [s or t * (avail - sum(short)) / sum(rest) for t, s in zip(text, short)]
        else:
            text = [t * avail / sum(text) for t in text]
        # no column narrower than its longest word; the shortfall comes out of the columns with room to spare
        floor = [max(len(w) for r in rows if j < len(r) for w in (r[j].split() or [""])) * CHAR_W for j in range(ncol)]
        short_by = sum(max(0.0, f - t) for f, t in zip(floor, text))
        room = sum(max(0.0, t - f) for f, t in zip(floor, text))
        if short_by and room > short_by:
            text = [f if t < f else t - (t - f) * short_by / room for f, t in zip(floor, text)]
        widths = [round((t + CELL_PAD) * 1440) for t in text]
        grid = tbl.find(_w("tblGrid"))
        for g in list(grid):
            grid.remove(g)
        for wdt in widths:
            grid.append(_set(OxmlElement("w:gridCol"), w=str(wdt)))
        tblpr = tbl.find(_w("tblPr"))
        for tag in ("tblW", "tblLayout"):
            old = tblpr.find(_w(tag))
            if old is not None:
                tblpr.remove(old)
        tblpr.find(_w("tblStyle")).addnext(_set(OxmlElement("w:tblW"), w=str(sum(widths)), type="dxa"))
        tblpr.append(_set(OxmlElement("w:tblLayout"), type="fixed"))
        _schema_order(tblpr, TBLPR_ORDER)
        for trpr in tbl.iter(_w("trPr")):
            _schema_order(trpr, TRPR_ORDER)
        trs = tbl.findall(_w("tr"))
        for i, tr in enumerate(trs):
            for tc, wdt in zip(tr.findall(_w("tc")), widths):
                tcpr = tc.find(_w("tcPr"))
                if tcpr is None:
                    tcpr = OxmlElement("w:tcPr")
                    tc.insert(0, tcpr)
                old = tcpr.find(_w("tcW"))
                if old is not None:
                    tcpr.remove(old)
                tcpr.insert(0, _set(OxmlElement("w:tcW"), w=str(wdt), type="dxa"))
                if i < len(trs) - 1:  # keep the table on one page
                    for p in tc.findall(_w("p")):
                        ppr = p.find(_w("pPr"))
                        if ppr is None:
                            ppr = OxmlElement("w:pPr")
                            p.insert(0, ppr)
                        if ppr.find(_w("keepNext")) is None:
                            ps = ppr.find(_w("pStyle"))
                            (ps.addnext if ps is not None else lambda x: ppr.insert(0, x))(OxmlElement("w:keepNext"))


def _sectpr(template, cols):
    """A continuous section break with `cols` columns, cloned from the body's page setup."""
    from docx.oxml import OxmlElement
    sect = deepcopy(template)
    for tag in ("type", "cols", "headerReference", "footerReference", "titlePg"):
        for el in sect.findall(_w(tag)):
            sect.remove(el)
    t = OxmlElement("w:type")
    t.set(_w("val"), "continuous")
    sect.insert(0, t)
    c = OxmlElement("w:cols")
    c.set(_w("num"), str(cols))
    c.set(_w("space"), str(round((TEXT_W - 2 * COL_W) * 1440)))
    sect.find(_w("pgMar")).addnext(c)
    return sect


def _carrier(sectpr):
    """An empty, zero-height paragraph whose only job is to end a section."""
    from docx.oxml import OxmlElement
    p, ppr = OxmlElement("w:p"), OxmlElement("w:pPr")
    ppr.append(_set(OxmlElement("w:spacing"), before="0", after="0", line="20", lineRule="exact"))
    ppr.append(sectpr)
    p.append(ppr)
    return p


def columns(doc, wide_tables: list[bool]) -> None:
    """Page setup, one-column title block, two-column body, and one-column sections around wide floats."""
    from docx.shared import Inches
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(TOP), Inches(BOTTOM)
    sec.left_margin = sec.right_margin = Inches(SIDE)
    sec.header_distance, sec.footer_distance = Inches(0.4), Inches(0.5)
    body = doc.element.body
    final = body.find(_w("sectPr"))
    for c in final.findall(_w("cols")):
        final.remove(c)
    final.find(_w("pgMar")).addnext(_sectpr(final, 2).find(_w("cols")))  # schema order: cols follows pgMar
    kids = list(body)
    first_h1 = next(i for i, el in enumerate(kids) if _style(el) == "Heading1")
    kids[first_h1].addprevious(_carrier(_sectpr(final, 1)))  # title block ends here

    col_emu = COL_W * 914400
    blocks, tables = [], iter(wide_tables)
    kids = list(body)
    for i, el in enumerate(kids):
        if i < first_h1:
            continue
        st = _style(el) if el.tag == _w("p") else None
        if st == "CaptionedFigure":
            ext = el.find(".//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent")
            if ext is not None and int(ext.get("cx")) > col_emu * 1.05:
                nxt = el.getnext()
                blocks.append((el, nxt if nxt is not None and _style(nxt) == "ImageCaption" else el))
        elif st == "TableCaption":
            tbl = el.getnext()
            while tbl is not None and tbl.tag != _w("tbl"):
                tbl = tbl.getnext()
            if tbl is not None and next(tables, False):
                blocks.append((el, tbl))

    def adjacent(a_end, b_start):  # only bookmarks / empty paragraphs in between
        el = a_end.getnext()
        while el is not None and el is not b_start:
            if el.tag == _w("tbl") or (el.tag == _w("p") and _text(el).strip()) or el.find(".//" + _w("drawing")) is not None:
                return False
            el = el.getnext()
        return el is b_start

    merged = []
    for start, end in blocks:
        if merged and adjacent(merged[-1][1], start):
            merged[-1] = (merged[-1][0], end)
        else:
            merged.append((start, end))
    for start, end in merged:
        start.addprevious(_carrier(_sectpr(final, 2)))
        end.addnext(_carrier(_sectpr(final, 1)))
    print(f"  {len(merged)} full-width regions")


def running_header(doc) -> None:
    """Right-aligned page number in the header, 8 pt (IEEE journal style)."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.shared import Pt
    p = doc.sections[0].header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fld = OxmlElement("w:fldSimple")
    fld.set(_w("instr"), "PAGE")
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(_w("val"), str(int(Pt(8).pt * 2)))
    rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    p._p.append(fld)


def keep_figures_together(doc) -> None:
    for p in doc.paragraphs:
        if p.style is not None and p.style.name in ("Captioned Figure", "Table Caption"):
            p.paragraph_format.keep_with_next = True


def build_docx() -> None:
    from docx import Document

    wide = flatten()
    try:
        restyle_reference_doc()
        convert()
    finally:
        FLAT.unlink(missing_ok=True)
    doc = Document(str(OUT_DOCX))
    front_matter(doc)
    number_and_crossref(doc)
    booktabs(doc)
    size_tables(doc, wide)
    keep_figures_together(doc)
    columns(doc, wide)
    running_header(doc)
    doc.save(str(OUT_DOCX))
    print(f"done: {OUT_DOCX}")


def main(argv) -> None:
    assert TEX_SRC.exists() and BIB.exists() and CSL.exists()
    want = {a.lstrip("-") for a in argv} or {"pdf", "docx"}
    if "pdf" in want:
        build_pdf()
    if "docx" in want:
        build_docx()


if __name__ == "__main__":
    main(sys.argv[1:])
